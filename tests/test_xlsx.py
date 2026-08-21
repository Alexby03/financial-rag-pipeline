"""Excel → Markdown tables and chunker breadcrumbs for financial models."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook

from app.config import Settings
from app.services.chunker import Chunker
from app.services.parser import DocumentParser
from app.services.xlsx import parse_xlsx_bytes


class _UnusedOcr:
    def process(self, file_bytes: bytes, filename: str):  # noqa: ANN001
        raise AssertionError("xlsx must not call OCR")


def _two_table_workbook_bytes() -> bytes:
    """Mirrors the P&L + FCF layout from the financial-model screenshot."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Model"

    # Table 1 — Total expenses breakdown / P&L
    sheet["A1"] = "Total expenses breakdown / P&L"
    headers = ["Financials (SEKm)", 1991, 1992, "1993E", "2003E", "2004E"]
    for col, value in enumerate(headers, start=1):
        sheet.cell(2, col, value)
    sheet["A3"] = "Total sales"
    sheet["B3"] = 100
    sheet["C3"] = 110
    sheet["D3"] = 120
    sheet["E3"] = 200
    sheet["F3"] = 210
    sheet["A4"] = "EBITDA"
    sheet["B4"] = 40
    sheet["C4"] = 44
    sheet["D4"] = 48
    sheet["E4"] = 80
    sheet["F4"] = 84
    sheet["A5"] = "Net income"
    sheet["B5"] = 20
    sheet["C5"] = 22
    sheet["D5"] = 24
    sheet["E5"] = 55.67
    sheet["F5"] = 58
    sheet["G4"] = "Opex TBU"
    sheet["G5"] = "orimligt att ta interest expense som % av sales"

    # Blank row gap → second visual table
    sheet["A7"] = "FCF from Net income"
    fcf_headers = ["SEKm", 1991, 1992, "1993E", "2003E", "2004E"]
    for col, value in enumerate(fcf_headers, start=1):
        sheet.cell(8, col, value)
    sheet["A9"] = "Net income"
    sheet["B9"] = 20
    sheet["C9"] = 22
    sheet["D9"] = 24
    sheet["E9"] = 55.67
    sheet["F9"] = 58
    sheet["A10"] = "Unlevered FCF"
    sheet["B10"] = 15
    sheet["C10"] = 16
    sheet["D10"] = 17
    sheet["E10"] = 40
    sheet["F10"] = 42

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_parse_xlsx_splits_blank_row_tables() -> None:
    parsed = parse_xlsx_bytes(_two_table_workbook_bytes())

    assert len(parsed.pages) == 1
    text = parsed.pages[0].text
    assert text.startswith("# Model")
    assert "## Total expenses breakdown / P&L" in text
    assert "## FCF from Net income" in text
    assert "| Post | 1991 | 1992 | 1993E | 2003E | 2004E |" in text or (
        "| Financials (SEKm) | 1991 | 1992 | 1993E | 2003E | 2004E |" in text
    )
    assert "55.67" in text
    assert "Unlevered FCF" in text
    assert "Opex TBU" in text
    assert parsed.annotation is not None
    assert parsed.annotation.year == 2004


def test_document_parser_bypasses_ocr_for_xlsx() -> None:
    parser = DocumentParser(ocr_client=_UnusedOcr())  # type: ignore[arg-type]
    parsed = parser.parse(_two_table_workbook_bytes(), "model.xlsx")
    assert len(parsed.pages) == 1
    assert "EBITDA" in parsed.pages[0].text


def test_chunker_keeps_each_xlsx_table_atomic_with_year_breadcrumbs() -> None:
    parsed = parse_xlsx_bytes(_two_table_workbook_bytes())
    chunker = Chunker(settings=Settings(openai_api_key="test-key", chunk_size=80, chunk_overlap=0))
    chunks = chunker.split_pages(parsed.pages)

    money = [chunk for chunk in chunks if "55.67" in chunk.text]
    fcf = [chunk for chunk in chunks if "Unlevered FCF" in chunk.text]

    assert money
    assert fcf
    assert money[0] is not fcf[0] or "Unlevered FCF" in money[0].text
    assert "| 1991 | 1992 | 1993E | 2003E | 2004E |" in money[0].text
    assert "Rows:" in money[0].text
    assert "Net income: 1991=20" in money[0].text or "Net income: 2004E=58" in money[0].text
    assert "Model > Total expenses breakdown / P&L" in money[0].text
    assert "Model > FCF from Net income" in fcf[0].text
    # Tables must not be character-split despite tiny chunk_size.
    assert "Net income" in money[0].text
    assert "1991" in money[0].text and "2004E" in money[0].text


def _side_by_side_workbook_bytes() -> bytes:
    """Two year-grids on the same rows with a blank column between them."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Parallel"

    # Left table
    sheet["A1"] = "Income statement"
    sheet["A2"] = "SEKm"
    sheet["B2"] = 2022
    sheet["C2"] = 2023
    sheet["A3"] = "Revenue"
    sheet["B3"] = 10
    sheet["C3"] = 12
    sheet["A4"] = "EBIT"
    sheet["B4"] = 2
    sheet["C4"] = 3

    # Blank column D, then right table
    sheet["F1"] = "Balance sheet"
    sheet["F2"] = "SEKm"
    sheet["G2"] = 2022
    sheet["H2"] = 2023
    sheet["F3"] = "Cash"
    sheet["G3"] = 5
    sheet["H3"] = 6
    sheet["F4"] = "Equity"
    sheet["G4"] = 8
    sheet["H4"] = 9

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_parse_xlsx_splits_side_by_side_tables() -> None:
    parsed = parse_xlsx_bytes(_side_by_side_workbook_bytes())
    text = parsed.pages[0].text

    assert "## Income statement" in text
    assert "## Balance sheet" in text
    assert "Revenue" in text and "Cash" in text
    # Must not smash both grids into one header row.
    assert text.count("| 2022 | 2023 |") >= 2
    income_pos = text.index("Income statement")
    balance_pos = text.index("Balance sheet")
    revenue_pos = text.index("Revenue")
    cash_pos = text.index("Cash")
    assert income_pos < revenue_pos < balance_pos or income_pos < balance_pos
    assert "Cash" in text[balance_pos:]
    assert "Revenue" in text[income_pos:balance_pos]


def test_side_by_side_notes_without_year_grid_stay_with_left_table() -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = "P&L"
    sheet["A2"] = "Post"
    sheet["B2"] = 2023
    sheet["C2"] = 2024
    sheet["A3"] = "Sales"
    sheet["B3"] = 1
    sheet["C3"] = 2
    # Blank column D, then a note-only strip (no year headers)
    sheet["F3"] = "check vs budget"
    buffer = BytesIO()
    workbook.save(buffer)

    text = parse_xlsx_bytes(buffer.getvalue()).pages[0].text
    assert text.count("## ") == 1
    assert "check vs budget" in text
    assert "Sales" in text


def test_side_by_side_banner_and_assumptions_stay_off_year_grid() -> None:
    """Regression: DCF models put General Assumptions beside the year grid on row 0."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "DCF"

    sheet["A1"] = "General Assumptions"
    sheet["G1"] = "Aberlyn (USDm)"
    sheet["H1"] = 1993
    sheet["I1"] = 1994
    sheet["G2"] = "Total revenue"
    sheet["H2"] = 1.0
    sheet["I2"] = 2.0
    sheet["G3"] = "Net income"
    sheet["H3"] = 0.5
    sheet["I3"] = 1.5

    # Banner is a few blank rows above the assumptions block (real Aberlyn layout).
    sheet["B5"] = "Case"
    sheet["C5"] = "Aberlyn"
    sheet["B6"] = "Terminal tax"
    sheet["C6"] = 0.34
    sheet["B7"] = "Sustained Growth Rate"
    sheet["C7"] = 0.02

    parsed = parse_xlsx_bytes(workbook_to_bytes(workbook))
    chunks = Chunker(Settings(openai_api_key="test-key")).split_pages(parsed.pages)

    revenue = next(chunk for chunk in chunks if "Total revenue" in chunk.text)
    assumptions = next(chunk for chunk in chunks if "Terminal tax" in chunk.text)

    assert "General Assumptions" not in revenue.text
    assert "Terminal tax" not in revenue.text
    assert "| 1993 | 1994 |" in revenue.text
    assert "General Assumptions" in assumptions.text
    assert assumptions.text.startswith("DCF > General Assumptions") or (
        "General Assumptions" in assumptions.text.split("\n\n", 1)[0]
    )
    assert "Sustained Growth Rate" in assumptions.text
    assert "| 1993 | 1994 |" not in assumptions.text
    # First data row must not become the section heading.
    assert not assumptions.text.split("\n\n", 1)[0].endswith("Terminal tax")
    assert " > Terminal tax" not in assumptions.text


def workbook_to_bytes(workbook: Workbook) -> bytes:
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_float_uses_excel_display_format() -> None:
    """Cached float 392.679 with format #,##0 must surface as 393, not long decimals."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = "Value"
    sheet["B1"] = 392.6792663085469
    sheet["B1"].number_format = "#,##0"
    sheet["A2"] = "Rate"
    sheet["B2"] = 0.017899
    sheet["B2"].number_format = "0.0%"

    from app.services.spreadsheet.loader import load_sheet_grids

    grid = next(iter(load_sheet_grids(workbook_to_bytes(workbook)).values()))
    assert grid[0][1] == "393"
    assert grid[1][1] == "1.8%"


def test_input_sheet_yes_no_stays_with_questions() -> None:
    """Blank columns before Yes/No must not create a lone Yes/No/No retrieval chunk."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Input Sheet"

    sheet["A1"] = "Adjustments to Current Financial Information"
    sheet["A2"] = "Do you have any operating leases?"
    sheet["D2"] = "Yes"
    sheet["A3"] = "Does your firm have R&D expenses?"
    sheet["D3"] = "No"
    sheet["A4"] = "Are there any other operating expenses to be capitalized?"
    sheet["D4"] = "No"

    text = parse_xlsx_bytes(workbook_to_bytes(workbook)).pages[0].text
    assert "Do you have any operating leases?" in text
    lease_line = next(line for line in text.splitlines() if "operating leases" in line)
    assert "Yes" in lease_line
    assert "| Yes |" not in text.split("operating leases")[0]  # no prior orphan Yes table


def test_label_value_gap_keeps_amounts_with_labels() -> None:
    """Blank column between label and amount must not orphan the number (Damodaran Options)."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Options"

    sheet["A26"] = "Value per option = "
    sheet["C26"] = 14.321702261167152
    sheet["C26"].number_format = '_("$"* #,##0.00_);_("$"* (#,##0.00);_("$"* "-"??_);_(@_)'
    sheet["A27"] = "Value of all options outstanding ="
    sheet["D27"] = 1150.605559662169
    sheet["D27"].number_format = '"$"#,##0_);[Red]\\("$"#,##0\\)'

    parsed = parse_xlsx_bytes(workbook_to_bytes(workbook))
    text = parsed.pages[0].text
    assert "Value per option" in text
    assert "14.32" in text
    assert "1151" in text
    # Same markdown table — not a lone-number Table N.
    option_line = next(line for line in text.splitlines() if "Value per option" in line)
    assert "14.32" in option_line


def test_options_dual_kv_columns_do_not_glue() -> None:
    """Left values must not glue onto the right-hand label column (touching neighbours)."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Options"

    sheet["A1"] = "VALUING WARRANTS WHEN THERE IS DILUTION"
    sheet["A2"] = "Stock Price="
    sheet["C2"] = 18.7386187
    sheet["D2"] = "# Warrants issued="
    sheet["F2"] = 80.34
    sheet["A3"] = "Strike Price="
    sheet["C3"] = 27.755
    sheet["D3"] = "# Shares outstanding="
    sheet["F3"] = 352

    text = parse_xlsx_bytes(workbook_to_bytes(workbook)).pages[0].text
    stock_line = next(line for line in text.splitlines() if "Stock Price" in line)
    warrants_line = next(line for line in text.splitlines() if "Warrants issued" in line)
    assert "18.738619" in stock_line or "18.74" in stock_line
    assert "80.34" in warrants_line
    # Left value must not lead the warrants row.
    assert not warrants_line.strip().startswith("| 18.")


def test_restated_financials_wide_gap_with_notes() -> None:
    """Labels several columns left of values + ! notes must rejoin on the same row."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Operating Leases"

    sheet["A1"] = "Restated Financials"
    sheet["A2"] = "Depreciation on Operating Lease Asset ="
    sheet["F2"] = 22.195710069381548
    sheet["F2"].number_format = "0.00"
    sheet["G2"] = "! I use straight line depreciation"
    sheet["A3"] = "Adjustment to Operating Earnings ="
    sheet["F3"] = 10.931387209170413
    sheet["F3"].number_format = "0.00"
    sheet["G3"] = "! PV of operating leases * Pre-tax cost of debt"
    sheet["A4"] = "Adjustment to Total Debt outstanding ="
    sheet["F4"] = 110.97855034690774
    sheet["F4"].number_format = "0.00"

    text = parse_xlsx_bytes(workbook_to_bytes(workbook)).pages[0].text
    dep = next(line for line in text.splitlines() if "Depreciation on Operating" in line)
    adj = next(line for line in text.splitlines() if "Adjustment to Operating Earnings" in line)
    debt = next(line for line in text.splitlines() if "Total Debt outstanding" in line)
    assert "22.2" in dep
    assert "straight line" in dep
    assert "10.93" in adj
    assert "110.98" in debt


def test_equity_value_sensitivity_stays_off_year_grid() -> None:
    """Right-hand sensitivity cube must not inherit DCF year-grid rows or title."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "DCF (50%)"

    sheet["G1"] = "Aberlyn (USDm)"
    sheet["H1"] = 1993
    sheet["I1"] = 1994
    sheet["G2"] = "Total revenue"
    sheet["H2"] = 0.876
    sheet["I2"] = 6.75
    sheet["G3"] = "Net income"
    sheet["H3"] = -0.7
    sheet["I3"] = 4.9
    # One blank row, then D&A continuation (real Aberlyn layout).
    sheet["G5"] = "Depreciation & Amortization"
    sheet["H5"] = 0.015
    sheet["I5"] = 0.12

    sheet["U1"] = "Equity Value Sensitivity"
    sheet["Y5"] = "Discount Rate:"
    # Scalar sits inside the cube bbox (Aberlyn: label + 12.8% two columns right).
    sheet["AA5"] = 0.128
    sheet["AA5"].number_format = "0.0%"
    sheet["Y6"] = "Revenue adj."
    sheet["Y7"] = -0.75
    sheet["Z7"] = -0.5
    sheet["AA7"] = 0
    sheet["Y9"] = 115.0
    sheet["Z9"] = 138.0
    sheet["AA9"] = 161.0
    sheet["V9"] = "Terminal growth"
    sheet["W9"] = 0.01
    sheet["W10"] = 0.02

    parsed = parse_xlsx_bytes(workbook_to_bytes(workbook))
    chunks = Chunker(Settings(openai_api_key="test-key")).split_pages(parsed.pages)
    equity = next(chunk for chunk in chunks if "Equity Value Sensitivity" in chunk.text)
    revenue = next(chunk for chunk in chunks if "Total revenue" in chunk.text)

    assert "Depreciation" not in equity.text
    assert "Total revenue" not in equity.text
    assert "Aberlyn (USDm)" not in equity.text
    assert "Discount Rate" in equity.text
    assert "12.8%" in equity.text
    assert " > Equity Value Sensitivity" in equity.text or equity.text.startswith(
        "DCF (50%) > Equity Value Sensitivity"
    )
    assert "Equity Value Sensitivity" not in revenue.text
    assert "Depreciation" in revenue.text


def test_chunker_does_not_leak_sheet_breadcrumbs_across_xlsx_pages() -> None:
    """Regression: Financials H1 must not stick onto Exhibit 7 antibody rows."""
    workbook = Workbook()
    financials = workbook.active
    assert financials is not None
    financials.title = "Financials"
    financials["A1"] = "P&L"
    financials["A2"] = "Post"
    financials["B2"] = 1991
    financials["C2"] = 1992
    financials["A3"] = "EBIT"
    financials["B3"] = 1
    financials["C3"] = 2

    exhibit = workbook.create_sheet("Exhibit 7 (Rhomed forecast)")
    exhibit["A1"] = "Antibody Derivative Products"
    exhibit["A2"] = "Post"
    exhibit["B2"] = 1991
    exhibit["C2"] = 1992
    exhibit["A3"] = "LeukoScan"
    exhibit["B3"] = "n.a."
    exhibit["C3"] = "n.a."

    buffer = BytesIO()
    workbook.save(buffer)
    parsed = parse_xlsx_bytes(buffer.getvalue())
    assert parsed.carry_page_context is False

    chunks = Chunker(Settings(openai_api_key="test-key")).split_pages(
        parsed.pages,
        carry_context=parsed.carry_page_context,
    )
    antibody = next(chunk for chunk in chunks if "LeukoScan" in chunk.text)
    assert antibody.page_number == 2
    assert "Exhibit 7 (Rhomed forecast) > Antibody Derivative Products" in antibody.text
    assert "Financials >" not in antibody.text
    assert "| Post | 1991 | 1992 |" in antibody.text
