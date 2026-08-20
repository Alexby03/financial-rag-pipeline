"""Load Excel workbooks into string grids per sheet."""

from __future__ import annotations

import re
from datetime import date, datetime
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

SheetGrids = dict[str, list[list[str]]]

# First section of an Excel number format before locale/positive;negative splits.
_FORMAT_SECTION = re.compile(r"^[^;]+")
_DECIMAL_ZEROS = re.compile(r"\.(0+)")
_DECIMAL_HASHES = re.compile(r"\.(#+)")


def load_sheet_grids(file_bytes: bytes) -> SheetGrids:
    """Return each worksheet as a trimmed string grid."""
    workbook = _load_workbook(file_bytes)
    try:
        grids: SheetGrids = {}
        for sheet in workbook.worksheets:
            grid = _sheet_grid(sheet)
            if grid:
                title = (sheet.title or "Sheet").strip() or "Sheet"
                grids[title] = grid
        return grids
    finally:
        workbook.close()


def _load_workbook(file_bytes: bytes) -> Workbook:
    workbook = load_workbook(BytesIO(file_bytes), data_only=True, read_only=False)
    if _workbook_has_values(workbook):
        return workbook
    workbook.close()
    return load_workbook(BytesIO(file_bytes), data_only=False, read_only=False)


def _workbook_has_values(workbook: Workbook) -> bool:
    for sheet in workbook.worksheets:
        max_row = min(sheet.max_row or 1, 40)
        max_col = min(sheet.max_column or 1, 40)
        for row in sheet.iter_rows(
            min_row=1,
            max_row=max_row,
            min_col=1,
            max_col=max_col,
            values_only=True,
        ):
            if any(value is not None and str(value).strip() for value in row):
                return True
    return False


def _sheet_grid(sheet: Worksheet) -> list[list[str]]:
    if not sheet.max_row or not sheet.max_column:
        return []

    grid: list[list[str]] = []
    for row in sheet.iter_rows(
        min_row=1,
        max_row=sheet.max_row,
        min_col=1,
        max_col=sheet.max_column,
    ):
        grid.append([_cell_to_text(cell) for cell in row])

    while grid and _row_blank(grid[-1]):
        grid.pop()
    if not grid:
        return []

    width = len(grid[0])
    while width > 0 and all(not row[width - 1].strip() for row in grid):
        for row in grid:
            row.pop()
        width -= 1
    return grid


def _row_blank(row: list[str]) -> bool:
    return all(not cell.strip() for cell in row)


def _cell_to_text(cell: Cell) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return _format_float(value, cell.number_format or "General")
    text = str(value).strip()
    if text.startswith("="):
        return ""
    return text


def _format_float(value: float, number_format: str) -> str:
    """Match Excel's *displayed* precision from ``number_format``, not raw IEEE bits.

    Example: cached value ``392.679…`` with format ``#,##0`` must become ``393``,
    not ``392.679266``.
    """
    if value.is_integer() and "%" not in number_format:
        return str(int(value))

    places = _decimal_places(number_format)
    is_percent = "%" in number_format
    display = value * 100.0 if is_percent else value

    if places is None:
        text = f"{display:.6f}".rstrip("0").rstrip(".")
        return f"{text}%" if is_percent else text
    if places == 0:
        rounded = str(int(round(display)))
        return f"{rounded}%" if is_percent else rounded
    rounded_value = round(display, places)
    text = f"{rounded_value:.{places}f}".rstrip("0").rstrip(".")
    return f"{text}%" if is_percent else text


def _decimal_places(number_format: str) -> int | None:
    """Return display decimals for a format, or None for General-style defaults."""
    match = _FORMAT_SECTION.match(number_format.strip())
    section = match.group(0) if match else number_format
    # Strip color/locale wrappers like [Red] or [$SEK]
    section = re.sub(r"\[[^\]]*]", "", section)
    if not section or section.lower() in {"general", "@"}:
        return None
    if "." not in section:
        # Formats like #,##0, 0, or 0% — round to integer for display.
        if any(token in section for token in ("0", "#")):
            return 0
        return None
    zeros = _DECIMAL_ZEROS.search(section)
    if zeros:
        return len(zeros.group(1))
    hashes = _DECIMAL_HASHES.search(section)
    if hashes:
        return len(hashes.group(1))
    return None
